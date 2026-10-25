"""Une el Bloque 1 redactado a mano (informe_fase1.docx) con el Bloque 2 de generar_informes.py.

Crea docs/informes/informe_entrega2.docx sin tocar informe_fase1.docx. El Bloque 2 sigue
la numeración del Bloque 1 (figuras desde la 8, tablas desde la 10) y sus referencias
empiezan en [9], porque [6] a [8] ya son la web y los notebooks. El índice y la
paginación se actualizan después en Word (ver docs/informes/README.md).
"""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import sys

from docx import Document
from docx.oxml.ns import qn

import generar_informes as g

BASE = g.DESTINO / 'informe_fase1.docx'
SALIDA = g.DESTINO / 'informe_entrega2.docx'
FECHA = '25/10/2026'
# Redacción del PDF de la entrega 1 (docs/entregas/PIDS_grupo2_entrega1.pdf) que difiere del .docx base.
FRASES = [
    ('Esto identifica un criterio por discutir, pero no establece un techo matemático del AP del modelo.',
     'Esto se debe a que la parte superior del robot es la más ambigua de anotar, debido a las antenas.'),
    ('No basta con equilibrar imágenes: V05 contiene 170 robots y V06 solo 117. V06 tiene 22 negativos en 93 imágenes, '
     'mientras V08 no tiene negativos y concentra 65 difficult. Se conservan estas diferencias porque representan '
     'situaciones reales de captura.',
     'No hemos intentado que todos los conjuntos sean exactamente iguales. Por ejemplo, V05 registra 170 robots, '
     'mientras que V06 contiene 117. Además, V06 incluye 22 ejemplos negativos, mientras que V08 no tiene ninguno y '
     'tiene 65 casos especialmente difíciles.'),
    ('Los canales RGB caracterizan el color dominante, no solo la iluminación, también dependen del suelo y las '
     'paredes. Por eso se complementan con media y desviación típica de intensidad, y con nitidez por varianza del '
     'Laplaciano. El tono azulado de exteriores justifica cambios de color suaves, así no se confunde ese tono con '
     'falta de luz.',
     'Los canales RGB reflejan el color predominante de cada imagen y no solo las condiciones de iluminación, ya que '
     'también están influidos por elementos como el suelo o las paredes. Por ello, se complementan con medidas de '
     'intensidad y nitidez. Además, el tono azulado habitual en exteriores justifica aplicar ajustes de color '
     'moderados para no confundirlo con una iluminación deficiente.'),
]
ESCENARIOS = {'V01': 'Vestíbulo ETSIT', 'V02': 'Aula A-201.L', 'V03': 'Aula A-201.L', 'V04': 'Exterior edificio A',
              'V05': 'Pasillo A-201.L', 'V06': 'Vestíbulo ETSIT (mesas)', 'V07': 'Vestíbulo ETSIT',
              'V08': 'Vestíbulo ETSIT (mesas)', 'V09': 'Exterior edificio A', 'V10': 'Aula A-201.L'}
NOTA_ANTIGUA = 'se han medido en validación, salvo las comprobaciones de integridad del aumento.'
NOTA_NUEVA = ('se han medido en validación, salvo las comprobaciones de integridad del aumento y la '
              'evaluación final, que se hizo en test una sola vez por modelo.')


def bloque2():
    """Devuelve (elementos del Bloque 2, documento de origen)."""
    inf = g.Informe('fase2')
    inf.cit = {tok: f'[{9 + i}]' for i, (tok, _) in enumerate(g.REFS_B2)}
    body = inf.doc.element.body
    toc = next(el for el in body if any(i.text and 'TOC' in i.text for i in el.iter(qn('w:instrText'))))
    inicio = list(body).index(toc) + 1
    inf.ntabla, inf.nfig = 9, 7
    inf.page()
    inf.fase2()
    return [e for e in list(body)[inicio:] if e.tag != qn('w:sectPr')], inf.doc


def copiar(elementos, origen, destino, ids):
    """Copia elementos al documento destino con sus imágenes y con ids de dibujo únicos."""
    copias = []
    for el in elementos:
        c = deepcopy(el)
        for blip in c.iter(qn('a:blip')):
            rid = blip.get(qn('r:embed'))
            blob = origen.part.related_parts[rid].blob
            nuevo, _ = destino.part.get_or_add_image(BytesIO(blob))
            blip.set(qn('r:embed'), nuevo)
        for ps in c.iter(qn('w:pStyle')):
            ps.set(qn('w:val'), destino.styles[origen.styles.element.get_by_id(ps.get(qn('w:val'))).name_val].style_id)
        for dp in c.iter('{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr'):
            dp.set('id', str(next(ids)))
        copias.append(c)
    return copias


def parrafo(doc, inicio):
    return next(p for p in doc.paragraphs if p.text.strip().startswith(inicio))


def poner_texto(p, texto):
    runs = p.runs
    if not runs:
        p.add_run(texto)
        return
    runs[0].text = texto
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def aplicar_pdf_entrega1(doc):
    """Lleva al Bloque 1 los cambios de redacción del PDF de la entrega 1."""
    for viejo, nuevo in FRASES:
        p = next(p for p in doc.paragraphs if viejo in p.text)
        poner_texto(p, p.text.replace(viejo, nuevo))
    # Estructura del PDF entregado: sin [7] ni [9] en el texto, apartado 2 como subtítulo y pies de figura con enlace.
    p = parrafo(doc, 'Tabla 2.')
    poner_texto(p, 'Tabla 2. Grabaciones propias y número exacto de imágenes extraídas.')
    next(p for p in doc.paragraphs if p.style.name == 'Heading 1' and p.text.strip() == '2. Recolección de datos').style = doc.styles['Heading 2']
    enlace = next(h for h in parrafo(doc, '[6]')._p.iter(qn('w:hyperlink')))
    for numero in (5, 6):
        p = parrafo(doc, f'Figura {numero}.')
        poner_texto(p, f'Figura {numero}. Gráficas del proyecto. Versión interactiva disponible en [6]: ')
        p._p.append(deepcopy(enlace))
    quitar = parrafo(doc, 'Las versiones interactivas de las Figuras 5 y 6')
    quitar._p.getparent().remove(quitar._p)
    quitar = parrafo(doc, '[9]')
    quitar._p.getparent().remove(quitar._p)
    tabla = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == 'Vídeo')
    for fila in tabla.rows[1:]:
        if fila.cells[0].text.strip() in ESCENARIOS:
            fila.cells[1].paragraphs[0].runs[0].text = ESCENARIOS[fila.cells[0].text.strip()]
            for r in fila.cells[1].paragraphs[0].runs[1:]:
                r.text = ''


def main():
    if len(sys.argv) == 1 and SALIDA.exists():
        print(f'{SALIDA.name} ya existe y es la versión final editada a mano en Word; '
              'regenerarla borraría esas ediciones. Pasa otra ruta de salida para generar una copia.')
        return
    doc = Document(BASE)
    aplicar_pdf_entrega1(doc)
    cuerpo, origen = bloque2()
    ids = iter(range(9000, 9999))

    celdas = doc.tables[0].rows
    celdas[2].cells[1].text = 'Entrega 2 · Bloques 1 y 2'
    celdas[3].cells[1].text = FECHA
    portada = parrafo(doc, 'MEMORIA DEL PROYECTO')
    poner_texto(portada, 'MEMORIA DEL PROYECTO  FASES 1 Y 2')

    referencias = next(p for p in doc.paragraphs if p.style.name == 'Heading 1' and p.text.strip() == 'Referencias')
    for el in copiar(cuerpo, origen, doc, ids):
        referencias._p.addprevious(el)

    ultima = parrafo(doc, '[8]')
    for numero, (_, texto) in enumerate(g.REFS_B2, start=9):
        copia = deepcopy(ultima._p)
        ultima._p.addnext(copia)
        ultima = type(ultima)(copia, ultima._parent)
        for h in list(copia.iter(qn('w:hyperlink'))):
            h.getparent().remove(h)
        poner_texto(ultima, f'[{numero}] {texto}')

    nota = parrafo(doc, 'Los originales, XML')
    assert NOTA_ANTIGUA in nota.text
    poner_texto(nota, nota.text.replace(NOTA_ANTIGUA, NOTA_NUEVA))

    doc.core_properties.title = 'Detección de RoboMasters Fases 1 y 2'
    salida = Path(sys.argv[1]) if len(sys.argv) > 1 else SALIDA
    doc.save(salida)
    print(salida)


if __name__ == '__main__':
    main()
